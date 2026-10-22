#!/usr/bin/env python3
"""Relit les sources officielles et régénère les données du site (RECETTE §4, §22).

Sources (toutes publiques, citées sur le site) :
  - KBA, Anschriftenverzeichnis der Zulassungsbehörden (AV 1), CSV + PDF (pour la date « Stand »)
  - KBA, Anschriftenverzeichnis der Fahrerlaubnisbehörden (AV 2), CSV
  - KBA, Statistik FZ 1 « Bestand an Kraftfahrzeugen und Kraftfahrzeuganhängern nach Zulassungsbezirken »
    (Datenlizenz Deutschland – Namensnennung 2.0)
  - Wikidata (CC0) : communes allemandes, clé communale (AGS), codes postaux, population

Usage :
  python3 scripts/update-data.py            # télécharge, puis écrit les JSON
  python3 scripts/update-data.py --offline  # rejoue le cache scripts/data/raw/
  python3 scripts/update-data.py --date 2026-10-22

Garde-fous : un fichier trop petit (téléchargement coupé) n'écrase rien ; la veille reste en place.
Dépendances : openpyxl (FZ 1), pypdf (date du PDF, facultatif).
"""
import csv, io, json, re, sys, datetime, urllib.request, urllib.parse, pathlib, collections

ROOT = pathlib.Path(__file__).resolve().parent.parent
RAW = ROOT / 'scripts' / 'data' / 'raw'
OUT_SRC = ROOT / 'src' / 'data' / 'zul'
OUT_PUB = ROOT / 'public' / 'data'
UA = {'User-Agent': 'Mozilla/5.0 (compatible; RadifPartners-data/1.0)'}
KBA = 'https://www.kba.de'
AV_PAGE = KBA + '/DE/Themen/ZentraleRegister/Anschriftenverzeichnisse/anschriftenverzeichnisse_node.html'
FZ1_PAGE = KBA + '/DE/Statistik/Fahrzeuge/Bestand/ZulassungsbezirkeGemeinden/zulassungsbezirke_node.html'
WIKIDATA = 'https://query.wikidata.org/sparql'
SPARQL = ('SELECT ?ags ?name (GROUP_CONCAT(DISTINCT ?plz;separator=" ") AS ?plzs) (SAMPLE(?pop) AS ?einw) WHERE { '
          '?g wdt:P439 ?ags. ?g wdt:P31/wdt:P279* wd:Q262166. FILTER NOT EXISTS { ?g wdt:P576 ?end } '
          'OPTIONAL{?g wdt:P281 ?plz} OPTIONAL{?g wdt:P1082 ?pop} ?g rdfs:label ?name. FILTER(LANG(?name)="de") } '
          'GROUP BY ?ags ?name')
LAND = {'01': 'SH', '02': 'HH', '03': 'NI', '04': 'HB', '05': 'NW', '06': 'HE', '07': 'RP', '08': 'BW', '09': 'BY',
        '10': 'SL', '11': 'BE', '12': 'BB', '13': 'MV', '14': 'SN', '15': 'ST', '16': 'TH'}


def get(url, timeout=120, headers=None, essais=3):
    """Téléchargement avec trois essais : le KBA et Wikidata répondent parfois 504."""
    import time
    for n in range(essais):
        try:
            req = urllib.request.Request(url, headers={**UA, **(headers or {})})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except Exception:
            if n == essais - 1:
                raise
            time.sleep(10 * (n + 1))


def link_on(page, pattern):
    html = get(page).decode('utf-8', 'ignore')
    m = re.findall(r'href="([^"]*' + pattern + r'[^"]*)"', html)
    if not m:
        raise SystemExit(f'lien introuvable ({pattern}) sur {page}')
    return urllib.parse.urljoin(page, m[0].replace('&amp;', '&'))


def fz1_url():
    """La page du KBA lie parfois une édition ancienne : on essaie l'année en cours puis les précédentes."""
    y = datetime.date.today().year
    for year in range(y, y - 4, -1):
        url = f'{KBA}/SharedDocs/Downloads/DE/Statistik/Fahrzeuge/FZ1/fz1_{year}.xlsx?__blob=publicationFile'
        try:
            if len(get(url)) > 100_000:
                return url
        except Exception:
            pass
    return link_on(FZ1_PAGE, r'fz1_[0-9]{4}\.xlsx')


def fetch_all():
    RAW.mkdir(parents=True, exist_ok=True)
    files = {
        'av1.csv': link_on(AV_PAGE, r'av1_[0-9_]+_csv\.csv'),
        'av1.pdf': link_on(AV_PAGE, r'av1_[0-9_]+_pdf\.pdf'),
        'av2.csv': link_on(AV_PAGE, r'av2_verzeichnis_Fahrerlaubnisbeh_csv\.csv'),
        'fz1.xlsx': fz1_url(),
    }
    mins = {'av1.csv': 100_000, 'av1.pdf': 100_000, 'av2.csv': 100_000, 'fz1.xlsx': 100_000}
    for name, url in files.items():
        data = get(url)
        if len(data) < mins[name]:
            raise SystemExit(f'{name} : {len(data)} octets, téléchargement suspect, rien n’est écrit')
        (RAW / name).write_bytes(data)
        print(f'{name:9} {len(data):>9} o  {url}')
    q = urllib.parse.urlencode({'query': SPARQL})
    data = get(f'{WIKIDATA}?{q}', timeout=180, headers={'Accept': 'text/csv'})
    if data.count(b'\n') < 10_000:
        raise SystemExit('Wikidata : moins de 10 000 communes, rien n’est écrit')
    (RAW / 'gemeinden.csv').write_bytes(data)
    print(f'gemeinden {data.count(b"\n"):>9} lignes')
    (RAW / 'urls.json').write_text(json.dumps(files, indent=1))


def av1_stand():
    try:
        import pypdf
        t = pypdf.PdfReader(str(RAW / 'av1.pdf')).pages[0].extract_text()
        m = re.search(r'Stand:\s*([A-Za-zäÄ]+ \d{4})', t)
        return m.group(1) if m else ''
    except Exception:
        return ''


# Kennzeichen von Bundes- und Sonderbehörden (Bundestag/Bundesregierung, Wasserstraßen, Bundespolizei, THW,
# diplomatische Kennzeichen, NATO, Bundeswehr) : sie stehen bei einzelnen Ämtern, gehören aber zu keinem Kreis.
SONDER = {'0', 'BD', 'BW', 'BP', 'BG', 'THW', 'X', 'Y',
          # Kennzeichen der Landesregierungen und Landtage (BWL, BYL, LSN…) : keine Kreiskennzeichen.
          'BBL', 'BWL', 'BYL', 'HEL', 'LSA', 'LSN', 'MVL', 'NL', 'NRW', 'RPL', 'SAL', 'SH', 'THL'}


def ok_uz(u):
    return bool(re.match(r'^[A-ZÄÖÜ]{1,3}$', u or '')) and u not in SONDER


def kreisname(behoerde, land):
    b = behoerde.strip()
    rest = b.split(' ', 1)[1] if ' ' in b else b
    # « LK Ostalbkreis », « LK Rhein-Kreis Neuss » : le nom porte déjà « Kreis », pas de préfixe.
    if b.startswith(('LK ', 'LRA ')) and 'kreis' in rest.lower():
        return rest
    if b.startswith('LK '):
        return ('Kreis ' if land in ('NW', 'SH') else 'Landkreis ') + b[3:]
    if b.startswith('LRA '):
        return 'Landkreis ' + b[4:]
    return b


def parse_av1():
    txt = (RAW / 'av1.csv').read_bytes().decode('latin-1')
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(txt), delimiter=';') if r]
    offices = collections.OrderedDict()
    for r in rows:
        key = r[0]
        if not re.match(r'^\d{5}-', key):
            continue
        name_parts = [x for x in r[3:8] if x]
        art = 'neben' if any('Nebenstelle' in x for x in name_parts) else 'haupt'
        behoerde = ' '.join(x for x in name_parts if x not in ('Nebenstelle',)).replace('- ', '-')
        if 'Polizei' in behoerde or 'Polizei' in r[1]:
            art = 'polizei'
        if 'Bundes' in behoerde or 'Bundes' in r[1]:
            art = 'bund'
        o = offices.setdefault(key, {
            'id': key.rstrip('-').replace('-e', 'e'), 'kz': key.split('-')[0],
            'zusatz': key.split('-')[1] if not key.split('-')[1].startswith('e') else '',
            'ort': r[1], 'land': r[2], 'behoerde': behoerde, 'art': art, 'plz': r[8], 'stadt': r[9],
            'strasse': r[10], 'tel': r[12].removeprefix('T:').strip(), 'mail': r[14], 'uz': [], 'alt': [], 'dv': r[25],
        })
        if ok_uz(r[15]) and r[15] not in o['uz']:
            o['uz'].append(r[15])
        for a in r[16:20]:
            if ok_uz(a) and a not in o['alt']:
                o['alt'].append(a)
    out = list(offices.values())
    for o in out:
        o['alt'] = [a for a in o['alt'] if a not in o['uz']]
    return out


def parse_av2():
    txt = (RAW / 'av2.csv').read_bytes().decode('latin-1')
    rows = list(csv.reader(io.StringIO(txt), delimiter=';'))[1:]
    recs = collections.OrderedDict()
    for r in rows:
        if len(r) < 8:
            continue
        key = r[6].strip()
        recs.setdefault(key, []).append([c.strip() for c in r])
    out = []
    for key, lines in recs.items():
        m = re.match(r'F (\d{5}) (\d{3})', key)
        if not m:
            continue
        first = lines[0]
        plz_ort = next((l[2] for l in lines[1:3] if re.match(r'^\d{5} ', l[2])), '')
        dept = ' · '.join(l[0] for l in lines[1:3] if l[0] and not l[0].startswith('F '))
        mail = next((l[3].replace('E-Mail:', '').strip() for l in lines if l[3].startswith('E-Mail')), '')
        tel = first[3].replace('Tel.:', '').strip()
        out.append({'kz': m.group(1), 'gem': m.group(2), 'name': first[0], 'stelle': dept, 'land': first[1],
                    'strasse': first[2], 'plz': plz_ort[:5], 'stadt': plz_ort[6:], 'tel': tel, 'mail': mail})
    return out


def parse_fz1():
    import openpyxl
    wb = openpyxl.load_workbook(str(RAW / 'fz1.xlsx'), read_only=True, data_only=True)
    imp = {r[1]: r[2] for r in wb['Impressum'].iter_rows(values_only=True) if r and len(r) > 2 and r[1]}
    stand = str(imp.get('Berichtszeitraum/Zeitpunkt:', ''))
    pub = str(imp.get('Veröffentlichung:', ''))
    s1 = {}
    for r in wb['FZ1.1'].iter_rows(values_only=True):
        if not r or not r[3] or not re.match(r'^\d{5} ', str(r[3])):
            continue
        k = str(r[3])[:5]
        s1[k] = {'name': str(r[3])[6:].strip(), 'krad': r[4] or 0, 'pkw': r[9] or 0, 'wohnmobile': r[16] or 0,
                 'halterinnen': r[19] or 0, 'pkwDichte': r[20] or 0, 'lkw': r[22] or 0, 'kfz': r[37] or 0,
                 'kfzDichte': r[38] or 0, 'anh': r[39] or 0}
    for r in wb['FZ1.2'].iter_rows(values_only=True):
        if not r or not r[3] or not re.match(r'^\d{5} ', str(r[3])):
            continue
        k = str(r[3])[:5]
        if k in s1:
            s1[k].update({'benzin': r[5] or 0, 'diesel': r[6] or 0, 'hybrid': r[8] or 0, 'phev': r[9] or 0, 'bev': r[10] or 0})
    tot = {}
    for r in wb['FZ1.1'].iter_rows(values_only=True):
        if r and r[1] == 'DEUTSCHLAND INSGESAMT':
            tot = {'krad': r[4], 'pkw': r[9], 'kfz': r[37], 'anh': r[39], 'pkwDichte': r[20]}
    for r in wb['FZ1.2'].iter_rows(values_only=True):
        if r and r[1] and str(r[1]).startswith('DEUTSCHLAND'):
            tot.update({'bev': r[10], 'phev': r[9], 'diesel': r[6], 'benzin': r[5]})
    return stand, pub, s1, tot


def parse_gemeinden():
    out = []
    for g in csv.DictReader(io.StringIO((RAW / 'gemeinden.csv').read_text('utf-8'))):
        ags = g['ags'].strip()
        if not re.match(r'^\d{8}$', ags):
            continue
        plz = sorted({p for p in g['plzs'].split() if re.match(r'^\d{5}$', p)})
        ein = int(float(g['einw'])) if g.get('einw') else 0
        out.append([g['name'], ags, plz, ein])
    # doublons (même AGS, deux libellés) : on garde le plus peuplé
    best = {}
    for x in out:
        if x[1] not in best or x[3] > best[x[1]][3]:
            best[x[1]] = x
    return sorted(best.values(), key=lambda x: (-x[3], x[0]))


def main():
    args = sys.argv[1:]
    if '--offline' not in args:
        fetch_all()
    build = args[args.index('--date') + 1] if '--date' in args else datetime.date.today().isoformat()
    offices = parse_av1()
    fe = parse_av2()
    stand, pub, stats, tot = parse_fz1()
    gem = parse_gemeinden()
    if len(offices) < 600 or len(stats) < 390 or len(gem) < 10_000 or len(fe) < 500:
        raise SystemExit(f'données incomplètes ({len(offices)} stellen, {len(stats)} bezirke, {len(gem)} gemeinden, {len(fe)} FE) : rien n’est écrit')
    # Zulassungsbezirke : un par clé de Kreis à 5 chiffres (+ villes à autorité propre « -1 »)
    bez = collections.OrderedDict()
    for o in offices:
        if o['art'] in ('polizei', 'bund'):
            continue
        k = o['kz'] + (('-' + o['zusatz']) if o['zusatz'] else '')
        b = bez.setdefault(k, {'id': k, 'kz': o['kz'], 'land': o['land'], 'name': '', 'uz': [], 'alt': [], 'stellen': []})
        if o['art'] == 'haupt' and not b['name']:
            b['name'] = o['ort']
            b['behoerde'] = o['behoerde']
        for u in o['uz']:
            if u not in b['uz']:
                b['uz'].append(u)
        for a in o['alt']:
            if a not in b['alt']:
                b['alt'].append(a)
        b['stellen'].append(o['id'])
    for b in bez.values():
        if not b['name']:
            first = next(o for o in offices if o['id'] == b['stellen'][0])
            b['name'] = first['ort']
            b['behoerde'] = first['behoerde']
        b['kreis'] = kreisname(b['behoerde'], b['land'])
        b['stat'] = stats.get(b['kz']) if not b['id'].count('-') else None
        b['fe'] = [i for i, f in enumerate(fe) if f['kz'] == b['kz']]
    # Unterscheidungszeichen -> Bezirke (courants et auslaufend)
    uz = collections.defaultdict(lambda: {'bez': [], 'alt': []})
    for b in bez.values():
        for u in b['uz']:
            uz[u]['bez'].append(b['id'])
        for a in b['alt']:
            uz[a]['alt'].append(b['id'])
    OUT_SRC.mkdir(parents=True, exist_ok=True)
    OUT_PUB.mkdir(parents=True, exist_ok=True)
    meta = {'build': build, 'av1Stand': av1_stand(), 'fz1Stand': stand, 'fz1Pub': pub,
            'stellen': len([o for o in offices if o['art'] in ('haupt', 'neben')]),
            'hauptstellen': len([o for o in offices if o['art'] == 'haupt']),
            'bezirke': len(bez), 'kennzeichen': len(uz), 'gemeinden': len(gem), 'fe': len(fe), 'total': tot}
    dump = lambda p, x: p.write_text(json.dumps(x, ensure_ascii=False, separators=(',', ':')) + '\n', 'utf-8')
    dump(OUT_SRC / 'fichier.json', meta)
    dump(OUT_SRC / 'bezirke.json', list(bez.values()))
    dump(OUT_SRC / 'kennzeichen.json', dict(sorted(uz.items())))
    dump(OUT_SRC / 'stellen.json', offices)
    dump(OUT_SRC / 'fe.json', fe)
    # Pour le navigateur : communes (nom, AGS, PLZ) chargées à la première recherche
    dump(OUT_PUB / 'gemeinden.json', [[g[0], g[1], g[2]] for g in gem])
    # Pour le navigateur : toutes les autorités (chargées à la première recherche, jamais au premier rendu)
    dump(OUT_PUB / 'behoerden.json', {'bezirke': list(bez.values()), 'stellen': offices, 'fe': fe, 'kennzeichen': dict(sorted(uz.items()))})
    print(json.dumps(meta, ensure_ascii=False))


if __name__ == '__main__':
    main()
